import openpyxl
import json

wb = openpyxl.load_workbook('Places/Sai_Gon_Places.xlsx')
sheet = wb['Sai_Gon']

headers = [cell.value for cell in sheet[1]]
rows = list(sheet.iter_rows(values_only=True))[1:]

stats = {
    "total": len(rows),
    "empty_open_hours": 0,
    "zero_price": 0,
    "null_photo": 0,
    "templated_description": 0,
    "unverified": 0,
    "unverified_items": [],
    "sample_items": []
}

for r in rows:
    p = dict(zip(headers, r))
    if p.get("open_hours") in ("{}", None, ""):
        stats["empty_open_hours"] += 1
    if p.get("price_vnd") in (0, None):
        stats["zero_price"] += 1
    if p.get("photo_url") is None or str(p.get("photo_url")).strip().lower() in ("", "nan", "none", "null"):
        stats["null_photo"] += 1
    desc = str(p.get("description", ""))
    if "thuộc nhóm" in desc and "cần dùng để gợi ý" in desc:
        stats["templated_description"] += 1
    if p.get("verified") is not True:
        stats["unverified"] += 1
        stats["unverified_items"].append({
            "name": p.get("name"),
            "kind": p.get("kind"),
            "lat": p.get("lat"),
            "lon": p.get("lon"),
            "note": p.get("note")
        })

print(f"Total: {stats['total']}")
print(f"Empty open_hours: {stats['empty_open_hours']}")
print(f"Zero price: {stats['zero_price']}")
print(f"Null photo: {stats['null_photo']}")
print(f"Templated description: {stats['templated_description']}")
print(f"Unverified: {stats['unverified']}")

checklist_sheet = wb['Checklist']
print("\nChecklist rows:")
for r in checklist_sheet.iter_rows(values_only=True):
    print(r)

sources_sheet = wb['Sources']
print("\nSources rows:")
for r in sources_sheet.iter_rows(values_only=True):
    print(r)

destinations_sheet = wb['Destinations']
print("\nDestinations rows:")
for r in destinations_sheet.iter_rows(values_only=True):
    print(r)

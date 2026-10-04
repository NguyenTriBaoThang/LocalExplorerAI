import openpyxl
import json
import sys

wb = openpyxl.load_workbook('Places/Sai_Gon_Places.xlsx')
report = {}

for sheetname in wb.sheetnames:
    sheet = wb[sheetname]
    headers = [cell.value for cell in sheet[1]]
    rows = list(sheet.iter_rows(values_only=True))
    total_data_rows = len(rows) - 1
    
    missing_by_col = {h: 0 for h in headers if h}
    for row in rows[1:]:
        for h, val in zip(headers, row):
            if not h:
                continue
            if val is None or str(val).strip() == "" or str(val).strip().lower() in ("none", "null", "nan"):
                missing_by_col[h] += 1
                
    report[sheetname] = {
        "total_rows": total_data_rows,
        "headers": headers,
        "missing_counts": missing_by_col,
    }

with open("scratch_excel_report.json", "w", encoding="utf-8") as f:
    json.dump(report, f, ensure_ascii=False, indent=2)

print("Report generated successfully.")

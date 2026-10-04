import openpyxl
import json

wb = openpyxl.load_workbook('Places/Sai_Gon_Places.xlsx')
data = {}
for sheetname in wb.sheetnames:
    sheet = wb[sheetname]
    headers = [cell.value for cell in sheet[1]]
    rows = []
    for r in list(sheet.iter_rows(values_only=True))[1:]:
        rows.append(dict(zip(headers, r)))
    data[sheetname] = rows

with open("scratch_all_sheets.json", "w", encoding="utf-8") as f:
    json.dump(data, f, ensure_ascii=False, indent=2)

print("Saved scratch_all_sheets.json")
